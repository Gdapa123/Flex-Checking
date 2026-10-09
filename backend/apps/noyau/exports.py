"""Exports Excel (openpyxl) et CSV (UTF-8 avec BOM, séparateur « ; » pour Excel en français)."""

import csv
import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import EmailEnvoye, ModeEmargement, StatutEmail
from .serializers import invites_annotes

LIBELLES_EMAIL = {StatutEmail.ENVOYE: "Envoyé", StatutEmail.EN_ATTENTE: "En attente", StatutEmail.ECHEC: "À renvoyer"}


def _dt(valeur):
    if not valeur:
        return ""
    return timezone.localtime(valeur).strftime("%d/%m/%Y %H:%M")


def _texte(valeur):
    """Neutralise une saisie libre interprétable comme formule par Excel (injection CSV/XLSX)."""
    if isinstance(valeur, str) and valeur[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + valeur
    return valeur


def _mode(code):
    return ModeEmargement(code).label if code else ""


def lignes_invites(qs):
    entetes = [
        "Événement",
        "Civilité",
        "Prénom",
        "Nom",
        "Email",
        "Téléphone",
        "Entreprise",
        "Fonction",
        "Statut",
        "Code",
        "Source",
        "Inscrit le",
        "Présent le",
        "Mode d’émargement",
        "Email 1",
        "Email 2",
    ]
    lignes = []
    for i in invites_annotes(qs).order_by("evenement__date", "nom", "prenom"):
        lignes.append(
            [
                i.evenement.nom_affiche,
                i.civilite_affichee,
                _texte(i.prenom),
                _texte(i.nom.upper()),
                _texte(i.email),
                i.telephone,
                _texte(i.entreprise),
                _texte(i.fonction),
                i.statut.libelle,
                i.code,
                i.get_source_display(),
                _dt(i.inscrit_le),
                _dt(i.present_le),
                _mode(i.mode_emargement),
                LIBELLES_EMAIL.get(i.email1, "—"),
                LIBELLES_EMAIL.get(i.email2, "—"),
            ]
        )
    return entetes, lignes


def lignes_presences(qs):
    entetes = [
        "Événement",
        "Civilité",
        "Prénom",
        "Nom",
        "Entreprise",
        "Présent le",
        "Mode d’émargement",
        "Poste",
        "Appareil",
        "Émargé par",
    ]
    lignes = []
    qs = (
        qs.filter(present_le__isnull=False)
        .select_related("evenement__programme")
        .prefetch_related("emargements__admin")
    )
    for i in qs.order_by("present_le"):
        em = next((e for e in i.emargements.all() if e.annule_le is None), None)
        lignes.append(
            [
                i.evenement.nom_affiche,
                i.civilite_affichee,
                _texte(i.prenom),
                _texte(i.nom.upper()),
                _texte(i.entreprise),
                _dt(i.present_le),
                _mode(i.mode_emargement),
                _texte(em.poste) if em else "",
                em.get_appareil_display() if em else "",
                em.admin.nom if em else "",
            ]
        )
    return entetes, lignes


def lignes_emails(qs_invites):
    entetes = ["Événement", "Invité", "Email", "Type", "Statut", "Tentatives", "Envoyé le", "Erreur"]
    lignes = []
    emails = EmailEnvoye.objects.filter(invite__in=qs_invites).select_related("invite__evenement__programme")
    for e in emails.order_by("cree_le"):
        lignes.append(
            [
                e.invite.evenement.nom_affiche,
                e.invite.nom_complet,
                _texte(e.invite.email),
                "Email n°1 · invitation" if e.type == "INVITATION" else "Email n°2 · magazine",
                e.get_statut_display(),
                e.tentatives,
                _dt(e.envoye_le),
                e.erreur[:300],
            ]
        )
    return entetes, lignes


def lignes_journal(entrees, cibles):
    from .journal import CATEGORIES
    from .journal_texte import decrire

    entetes = ["Date", "Utilisateur", "Type", "Détail"]
    lignes = [
        [_dt(e.date), e.auteur.nom if e.auteur else "Système", CATEGORIES.get(e.action, ""), decrire(e, cibles)]
        for e in entrees
    ]
    return entetes, lignes


def en_csv(entetes, lignes):
    tampon = io.StringIO()
    tampon.write("﻿")
    writer = csv.writer(tampon, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(entetes)
    writer.writerows(lignes)
    return tampon.getvalue().encode("utf-8")


def en_xlsx(entetes, lignes, titre="Export"):
    wb = Workbook()
    ws = wb.active
    ws.title = titre[:31]
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1B1D20")
        cellule.alignment = Alignment(vertical="center")
    for ligne in lignes:
        ws.append(ligne)
    for idx, entete in enumerate(entetes, start=1):
        largeur = max([len(str(entete))] + [len(str(r[idx - 1])) for r in lignes[:500]]) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(largeur, 48)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def fichier(entetes, lignes, fmt, titre):
    if fmt == "csv":
        return en_csv(entetes, lignes), "text/csv; charset=utf-8", "csv"
    return (
        en_xlsx(entetes, lignes, titre),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    )
