import pytest
from django.core import mail

from apps.noyau.models import EmailEnvoye, Emargement, Invite, JournalActivite, StatutEmail

from .conftest import nouvel_invite

pytestmark = pytest.mark.django_db


def url(ev, s=""):
    return f"/api/evenements/{ev.pk}/invites/{s}"


@pytest.fixture
def peuple(evenement):
    a = nouvel_invite(evenement, email="aya@sib.ci", nom="Kouassi", entreprise="SIB")
    b = nouvel_invite(evenement, email="yao@orange.ci", nom="Konan", prenom="Yao", civilite="M", entreprise="Orange")
    c = nouvel_invite(evenement, email="awa@gmail.com", nom="Traoré", prenom="Awa", telephone="+225 05 11 22 33 44")
    return a, b, c


def test_liste_compteurs_pagination(client_admin, evenement, peuple):
    client_admin.post(f"/api/invites/{peuple[0].pk}/presence/")
    d = client_admin.get(url(evenement), {"taille": 25}).json()
    assert d["total"] == 3 and d["taille"] == 25
    assert d["compteurs"] == {"tous": 3, "inscrits": 2, "presents": 1, "aRenvoyer": 0}
    assert {"email1", "email2", "statut", "nomComplet", "code"} <= set(d["resultats"][0])


@pytest.mark.parametrize("q,attendu", [("kouassi", 1), ("orange", 1), ("05 11 22", 1), ("gmail", 1), ("zzz", 0)])
def test_recherche(client_admin, evenement, peuple, q, attendu):
    assert client_admin.get(url(evenement), {"q": q}).json()["total"] == attendu


def test_recherche_par_code(client_admin, evenement, peuple):
    assert client_admin.get(url(evenement), {"q": peuple[1].code}).json()["resultats"][0]["id"] == peuple[1].pk


def test_filtres_et_tri(client_admin, evenement, peuple):
    client_admin.post(f"/api/invites/{peuple[2].pk}/presence/")
    d = client_admin.get(url(evenement), {"statut": "PRESENT"}).json()
    assert [r["id"] for r in d["resultats"]] == [peuple[2].pk]
    assert d["resultats"][0]["modeEmargement"] == "MANUEL"
    assert client_admin.get(url(evenement), {"mode": "MANUEL"}).json()["total"] == 1
    assert client_admin.get(url(evenement), {"magazine": "ENVOYE"}).json()["total"] == 1
    assert client_admin.get(url(evenement), {"magazine": "NON"}).json()["total"] == 2
    noms = [r["nom"] for r in client_admin.get(url(evenement), {"tri": "nom"}).json()["resultats"]]
    assert noms == sorted(noms)
    EmailEnvoye.objects.filter(invite=peuple[1]).update(statut=StatutEmail.ECHEC)
    assert client_admin.get(url(evenement), {"email1": "ECHEC"}).json()["total"] == 1


@pytest.mark.parametrize(
    "tri,attendu",
    [
        ("nom", ["Konan", "Kouassi", "Traoré"]),
        ("-nom", ["Traoré", "Kouassi", "Konan"]),
        ("prenom", ["Traoré", "Kouassi", "Konan"]),  # Awa, Aya, Yao
        ("-prenom", ["Konan", "Kouassi", "Traoré"]),
        ("entreprise", ["Konan", "Kouassi", "Traoré"]),  # Orange, puis SIB par nom
        ("-entreprise", ["Kouassi", "Traoré", "Konan"]),  # SIB par nom (A → Z), puis Orange
        ("inscritLe", ["Kouassi", "Konan", "Traoré"]),
        ("-inscritLe", ["Traoré", "Konan", "Kouassi"]),
        ("-statut", ["Konan", "Kouassi", "Traoré"]),  # présent d'abord, puis par nom
        ("-presentLe", ["Konan", "Kouassi", "Traoré"]),  # arrivée la plus récente, sans arrivée en dernier
    ],
)
def test_options_de_tri(client_admin, evenement, peuple, tri, attendu):
    client_admin.post(f"/api/invites/{peuple[1].pk}/presence/")
    noms = [r["nom"] for r in client_admin.get(url(evenement), {"tri": tri}).json()["resultats"]]
    assert noms == attendu


def test_ajout_manuel(client_admin, evenement):
    mail.outbox.clear()
    r = client_admin.post(
        url(evenement),
        {
            "civilite": "M",
            "nom": "Bamba",
            "prenom": "Franck",
            "email": "franck@sib.ci",
            "telephone": "0707070707",
            "entreprise": "SIB",
            "fonction": "Auditeur",
            "envoyerInvitation": False,
        },
        format="json",
    )
    assert r.status_code == 201, r.content
    inv = Invite.objects.get(email="franck@sib.ci")
    assert inv.source == "MANUEL" and len(mail.outbox) == 0
    assert JournalActivite.objects.filter(action="INVITE_AJOUT", cible_id=inv.pk).exists()
    r = client_admin.post(
        url(evenement),
        {"civilite": "M", "nom": "Bamba", "prenom": "F", "email": "franck@sib.ci", "telephone": "0707070707"},
        format="json",
    )
    assert r.status_code in (400, 409)


def test_ajout_manuel_exige_entreprise_et_fonction(client_admin, evenement):
    r = client_admin.post(
        url(evenement),
        {"civilite": "M", "nom": "Bamba", "prenom": "Franck", "email": "f2@sib.ci", "telephone": "0707070707"},
        format="json",
    )
    assert r.status_code == 400
    assert r.json()["entreprise"] == ["Indiquez votre entreprise."]
    assert r.json()["fonction"] == ["Indiquez votre fonction."]


def test_modification_et_historique(client_admin, evenement, peuple):
    inv = peuple[0]
    r = client_admin.patch(f"/api/invites/{inv.pk}/", {"entreprise": "NSIA Banque"}, format="json")
    assert r.status_code == 200 and r.json()["entreprise"] == "NSIA Banque"
    client_admin.post(f"/api/invites/{inv.pk}/presence/")
    client_admin.post(f"/api/invites/{inv.pk}/presence/annulation/", {"motif": "Test"}, format="json")
    types = [e["type"] for e in client_admin.get(f"/api/invites/{inv.pk}/historique/").json()]
    assert types[0] == "INSCRIPTION"
    assert {"EMAIL_ENVOYE", "EMARGEMENT", "ANNULATION", "MODIFICATION"} <= set(types)


def test_renvoi_individuel(client_admin, peuple):
    mail.outbox.clear()
    r = client_admin.post(f"/api/invites/{peuple[0].pk}/renvoi-email/", {}, format="json")
    assert r.status_code == 202 and len(mail.outbox) == 1


def test_actions_groupees_synchrones(client_admin, evenement, peuple):
    ids = [p.pk for p in peuple]
    r = client_admin.post(url(evenement, "actions-groupees/"), {"action": "VALIDER", "ids": ids}, format="json")
    assert r.status_code == 200 and r.json()["traites"] == 3
    assert Emargement.objects.filter(invite__in=ids, mode="MANUEL").count() == 3
    r = client_admin.post(url(evenement, "actions-groupees/"), {"action": "RENVOYER", "ids": ids[:2]}, format="json")
    assert r.json()["traites"] == 2


def test_actions_groupees_celery_au_dela_de_50(client_admin, evenement):
    for i in range(55):
        nouvel_invite(evenement, email=f"g{i}@sib.ci")
    r = client_admin.post(
        url(evenement, "actions-groupees/") + "?statut=INSCRIT", {"action": "VALIDER", "tous": True}, format="json"
    )
    assert r.status_code == 202
    tache = client_admin.get(f"/api/taches/{r.json()['tacheId']}/").json()
    assert tache["statut"] == "TERMINEE" and tache["traites"] == 55


def test_export_csv_bom_et_xlsx(client_admin, evenement, peuple):
    r = client_admin.get(url(evenement, "export/"), {"format": "csv"})
    assert r.status_code == 200
    contenu = r.content
    assert contenu.startswith(b"\xef\xbb\xbf")
    assert "Kouassi".upper() in contenu.decode("utf-8-sig")
    r = client_admin.get(url(evenement, "export/"), {"format": "xlsx", "statut": "INSCRIT"})
    assert r["Content-Type"].startswith("application/vnd.openxmlformats")
    assert JournalActivite.objects.filter(action="EXPORT").count() == 2


def test_export_selection_async(client_admin, evenement):
    ids = [nouvel_invite(evenement, email=f"e{i}@sib.ci").pk for i in range(60)]
    r = client_admin.post(
        url(evenement, "actions-groupees/"), {"action": "EXPORTER", "ids": ids, "format": "xlsx"}, format="json"
    )
    assert r.status_code == 202
    t = client_admin.get(f"/api/taches/{r.json()['tacheId']}/").json()
    assert t["statut"] == "TERMINEE" and t["nb"] == 60
    f = client_admin.get(t["telechargement"])
    assert f.status_code == 200
    # fichier à usage unique : supprimé après téléchargement
    assert client_admin.get(t["telechargement"]).status_code == 404


def test_qr_png(client_admin, peuple):
    r = client_admin.get(f"/api/invites/{peuple[0].pk}/qr.png")
    assert r.status_code == 200 and r.content[:4] == b"\x89PNG"


def test_statistiques(client_admin, evenement, peuple):
    client_admin.post(f"/api/invites/{peuple[0].pk}/presence/")
    d = client_admin.get(f"/api/evenements/{evenement.pk}/statistiques/").json()
    assert d["inscrits"] == 3 and d["presents"] == 1 and d["taux"] == 33
    assert sum(m["nb"] for m in d["modes"]) == 1
    assert d["comparaison"][0]["enCours"] is True


def test_tranche_en_cours_calculee_par_le_serveur(client_admin, evenement):
    from datetime import timedelta

    from django.utils import timezone

    ici = timezone.localtime()
    if ici.hour < 1 or ici.hour >= 22:
        pytest.skip("cérémonie du jour impossible à placer autour de minuit")
    evenement.date = ici.date()
    evenement.date_limite_inscription = ici.date()
    evenement.heure_accueil = (ici - timedelta(minutes=40)).time()
    evenement.heure_debut = (ici + timedelta(minutes=20)).time()
    evenement.save()
    d = client_admin.get(f"/api/evenements/{evenement.pk}/statistiques/").json()
    tranche = d["tranches"][d["trancheEnCours"]]
    h, m = (int(x) for x in tranche["debut"].split("h"))
    assert h * 60 + m <= ici.hour * 60 + ici.minute < h * 60 + m + 15
    # autre jour que la cérémonie : pas de tranche en cours
    evenement.date = evenement.date_limite_inscription = ici.date() + timedelta(days=5)
    evenement.save()
    assert client_admin.get(f"/api/evenements/{evenement.pk}/statistiques/").json()["trancheEnCours"] is None


def test_export_neutralise_les_formules(client_admin, evenement):
    """Une valeur saisie par un invité ne doit jamais être interprétée comme une formule (CSV/Excel)."""
    import io

    from openpyxl import load_workbook

    nouvel_invite(evenement, email="f@sib.ci", entreprise='=HYPERLINK("http://evil","Clic")', prenom="+Aya")
    csv = client_admin.get(url(evenement, "export/"), {"format": "csv"}).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv and ";=HYPERLINK" not in csv
    xlsx = client_admin.get(url(evenement, "export/"), {"format": "xlsx"}).content
    valeurs = [c.value for row in load_workbook(io.BytesIO(xlsx)).active.iter_rows() for c in row]
    assert '\'=HYPERLINK("http://evil","Clic")' in valeurs
    assert not any(isinstance(v, str) and v.startswith("=") for v in valeurs)


def test_purge_des_exports_anciens(tmp_path, monkeypatch):
    import os
    import time

    from apps.noyau import taches

    monkeypatch.setattr(taches, "DOSSIER", tmp_path)
    vieux, recent = tmp_path / "vieux.xlsx", tmp_path / "recent.xlsx"
    vieux.write_bytes(b"x")
    recent.write_bytes(b"x")
    il_y_a_7h = time.time() - 7 * 3600
    os.utime(vieux, (il_y_a_7h, il_y_a_7h))
    assert taches.purger_exports() == 1
    assert not vieux.exists() and recent.exists()
